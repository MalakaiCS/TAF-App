"""Getting a worksheet onto paper on a machine that is not Windows.

The worksheet is the one piece of this that has to exist physically. It goes
on the bench next to the saw, and somebody works from it all morning. For
years the only way to make it was Excel driving itself through COM, which is
Windows and nothing else - so on a Mac the app saved a spreadsheet, opened it,
and left the rest to whoever was standing there.

These are about the converters underneath: that each platform tries the thing
most likely to be installed on it first, that a converter which cannot work
here says so instead of pretending, and that when every one of them fails the
order is still saved rather than lost.
"""
from __future__ import annotations

import os
import platform
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import template_filler as _tf          # noqa: E402

SRC = (ROOT / "template_filler.py").read_text(encoding="utf-8")


# ── Actually making one ──────────────────────────────────────────────────────

def test_libreoffice_really_turns_the_template_into_a_pdf():
    """Not a mock. The template is loaded, saved the way the app saves it, and
    converted - because the failure this is guarding against was LibreOffice
    being present with no spreadsheet filters installed, which no amount of
    stubbing would have caught."""
    if not _tf._soffice():
        return                      # nothing to test on a machine without it
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "O-N-9999.xlsx")
        wb = openpyxl.load_workbook(ROOT / "Templates.xlsx")
        made = _tf.save_and_export_pdf(wb, out, auto_open=False)
        assert made.endswith(".pdf"), f"it gave back {made}"
        assert os.path.exists(made)
        with open(made, "rb") as fh:
            assert fh.read(5) == b"%PDF-", "that file is not a PDF"
        assert os.path.getsize(made) > 2000, "a PDF that small has no sheet in it"


def test_the_spreadsheet_is_saved_before_anything_is_converted():
    """Every converter can fail. None of them may take the order with it."""
    body = SRC.split("def save_and_export_pdf")[1].split("\ndef ")[0]
    save_at = body.index("out_wb.save(out_path)")
    assert save_at < body.index("ways ="), \
        "it decides how to convert before the order is written down"


def test_nothing_to_convert_with_still_leaves_the_order_on_disk():
    import openpyxl
    was = _tf._soffice
    _tf._soffice = lambda: ""           # no LibreOffice, and this is Linux
    try:
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "O-N-9998.xlsx")
            wb = openpyxl.Workbook()
            wb.active["A1"] = "O/N 9998"
            made = _tf.save_and_export_pdf(wb, out, auto_open=False)
            assert made == out, "it claimed a PDF it never made"
            assert os.path.exists(out), "the order itself is gone"
    finally:
        _tf._soffice = was


def test_a_converter_that_cannot_run_here_says_no_rather_than_raising():
    """Excel for Mac on a Linux box. It has to answer False, not throw, or
    the loop over the converters never reaches LibreOffice."""
    if platform.system() == "Darwin":
        return
    assert _tf._export_pdf_excel_mac("/nowhere/x.xlsx", "/nowhere/x.pdf") is False


def test_a_missing_libreoffice_is_a_no_not_a_crash():
    was = _tf._soffice
    _tf._soffice = lambda: ""
    try:
        assert _tf._export_pdf_libreoffice("/nowhere/x.xlsx", "/nowhere/x.pdf") is False
    finally:
        _tf._soffice = was


def test_it_does_not_claim_success_off_a_returncode_alone():
    """soffice exits 0 having written nothing rather often - a locked profile,
    a filter it does not have. Believing the exit code means the app says the
    worksheet printed and there is no PDF."""
    body = SRC.split("def _export_pdf_libreoffice")[1].split("\ndef ")[0]
    assert "os.path.exists(made)" in body
    assert body.rstrip().endswith("return os.path.exists(pdf_path)")


# ── Which one is tried first ─────────────────────────────────────────────────

def _ways(system: str) -> str:
    body = SRC.split("def save_and_export_pdf")[1].split("\ndef ")[0]
    if system == "Windows":
        return body.split('if system == "Windows":')[1].split("elif")[0]
    if system == "Darwin":
        return body.split('elif system == "Darwin":')[1].split("else:")[0]
    return body.split("else:\n        ways =")[1].split("\n\n")[0]


def test_windows_still_reaches_for_excel_first():
    """Nothing about the other two platforms may change what the factory's own
    machines do. Excel matches the template exactly; it goes first."""
    order = _ways("Windows")
    assert order.index("cached_excel") < order.index("powershell") < \
        order.index("libreoffice")


def test_a_mac_tries_excel_before_libreoffice():
    """The template is a fixed grid read at a saw. LibreOffice moves a column
    here and there, which is fine as a fallback and wrong as a first choice
    when the real thing is installed."""
    order = _ways("Darwin")
    assert order.index("excel_mac") < order.index("libreoffice")


def test_a_mac_never_tries_com():
    """win32com does not exist there. Reaching for it is an import error in
    the middle of making an order."""
    order = _ways("Darwin")
    assert "cached_excel" not in order and "powershell" not in order


def test_every_platform_has_something_to_try():
    for system in ("Windows", "Darwin", "Linux"):
        assert "libreoffice" in _ways(system), system


# ── Opening it afterwards ────────────────────────────────────────────────────

def test_opening_a_file_knows_all_three_platforms():
    body = SRC.split("def _open_file")[1].split("\ndef ")[0]
    assert "os.startfile" in body
    assert '"open"' in body
    assert '"xdg-open"' in body


def test_nothing_is_opened_when_nobody_asked():
    """Bulk printing calls this in a loop. Fifty PDFs opening on top of each
    other is how somebody force-quits halfway through a run."""
    import openpyxl
    opened: list[str] = []
    was = _tf._open_file
    _tf._open_file = lambda path: opened.append(path)
    try:
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(ROOT / "Templates.xlsx")
            _tf.save_and_export_pdf(wb, os.path.join(tmp, "O-N-9997.xlsx"),
                                    auto_open=False)
        assert opened == [], f"it opened {opened}"
    finally:
        _tf._open_file = was
